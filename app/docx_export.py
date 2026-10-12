"""Сборка Word-документа из распознанного текста (Markdown-разметка модели)."""
import io
import re
from dataclasses import dataclass

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.shared import Cm, Pt, RGBColor


@dataclass
class Page:
    text: str
    landscape: bool = False


# Разделитель шапки Markdown-таблицы: |---|:---:|
_SEPARATOR_CELL = re.compile(r"^:?-{3,}:?$")
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _is_table_row(line: str) -> bool:
    return line.startswith("|") or line.count("|") >= 2


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def _is_separator(cells: list[str]) -> bool:
    return all(_SEPARATOR_CELL.match(c) for c in cells if c) and any(cells)


def _blocks(text: str):
    """Разбивает текст на блоки: ("heading", level, text), ("table", rows, header), ("para", text)."""
    lines = [line.strip() for line in text.splitlines()]
    # Модель иногда ставит «#» перед строкой таблицы — это строка таблицы, а не заголовок
    lines = [line.lstrip("#").strip() if line.startswith("#") and line.count("|") >= 2 else line for line in lines]
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line:
            i += 1
            continue
        if m := _HEADING.match(line):
            yield "heading", len(m.group(1)), m.group(2)
            i += 1
        elif _is_table_row(line):
            rows, header = [], False
            while i < len(lines) and lines[i] and _is_table_row(lines[i]):
                cells = _split_row(lines[i])
                if _is_separator(cells):
                    header = header or len(rows) == 1
                else:
                    rows.append(cells)
                i += 1
            if rows:
                yield "table", rows, header
        else:
            yield "para", line
            i += 1


def _merge_tables(blocks: list) -> list:
    """Склеивает соседние таблицы с одинаковым числом столбцов (таблица, разрезанная на полосы)."""
    merged = []
    for block in blocks:
        prev = merged[-1] if merged else None
        if (
            block[0] == "table" and prev and prev[0] == "table"
            and len(block[1][0]) == len(prev[1][0])
        ):
            rows = block[1]
            # Повтор шапки в продолжении таблицы не дублируем
            if block[2] and rows[0] == prev[1][0]:
                rows = rows[1:]
            merged[-1] = ("table", prev[1] + rows, prev[2])
        else:
            merged.append(block)
    return merged


def _add_text(paragraph, text: str, bold: bool = False):
    """Добавляет текст с поддержкой **жирного** и <br>."""
    parts = text.replace("<br/>", "<br>").replace("<br />", "<br>").split("<br>")
    for n, part in enumerate(parts):
        if n:
            paragraph.add_run().add_break()
        pos = 0
        for m in _BOLD.finditer(part):
            if m.start() > pos:
                paragraph.add_run(part[pos:m.start()]).bold = bold or None
            paragraph.add_run(m.group(1)).bold = True
            pos = m.end()
        if pos < len(part):
            paragraph.add_run(part[pos:]).bold = bold or None


def _add_table(doc: Document, rows: list[list[str]], header: bool):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r, cells in enumerate(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            _add_text(cell.paragraphs[0], cells[c] if c < len(cells) else "", bold=header and r == 0)


def _set_orientation(section, landscape: bool):
    width, height = Cm(21), Cm(29.7)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (height, width) if landscape else (width, height)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))


def build_docx(pages: list[Page]) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(2)
    for level, size in ((1, 14), (2, 13), (3, 12)):
        heading = doc.styles[f"Heading {level}"]
        heading.font.name = "Times New Roman"
        heading.font.size = Pt(size)
        heading.font.color.rgb = RGBColor(0, 0, 0)

    for n, page in enumerate(pages):
        # Каждая страница PDF — новый раздел Word со своей ориентацией
        section = doc.sections[0] if n == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
        _set_orientation(section, page.landscape)
        for block in _merge_tables(list(_blocks(page.text))):
            if block[0] == "heading":
                _add_text(doc.add_heading(level=min(block[1], 3)), block[2])
            elif block[0] == "table":
                _add_table(doc, block[1], block[2])
                doc.add_paragraph()
            else:
                _add_text(doc.add_paragraph(), block[1])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
